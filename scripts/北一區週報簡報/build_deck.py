#!/usr/bin/env python3
"""北一區週報簡報（Keynote）自動產生。

    python3 build_deck.py                  # 上週（最近完整的週日~週六）
    python3 build_deck.py --wk-end 2026-10-10 --out-dir /tmp/x   # 指定週、輸出到別處（測試用）

流程：複製上一週的 .key → 撈各來源數字 → 寫入表格 / 換圖 / 改週次標題 → 存檔。
評語（上週店點狀況／執行進度與改善狀況）不在這裡寫：本程式產出「評語素材.md」，
由排程中的 Claude 依素材寫好 評語.json 後，再跑 fill_comments.py 填入並匯出 PDF。

資料來源（逐格已對 27Q1W02 手工版驗證，見 README.md）：
  第2頁 業績總表      dss@ 每週日寄的「北一區門市週報_YYYYMMDD」附件（照它的公式重算）
                      ＋ 來客數（週報產生器的 ShopperTrak 人流）＋ 銷售人數（排班系統）
  重點項目進度        週報產生器 Excel「重點項目進度」；Mac mini M4 直查 EPB
  MySetup             Apple Personal Setup 信件（warranty_report_all.scan_mysetup）＋週報 Excel 服務費
  3PP 搭售統計        dss_rebuild_from_epb（EPB 重建 DSS 3PP）
  維修 營收／圖表      repair_report（EPB）
  手填（黃底）         維修預約、SACare 整機搭售、CSAT —— 來源是週二才寄的維修週報
"""
import argparse
import calendar
import email
import glob
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from email import policy
from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / '排班系統'))
MAIL_DB = Path.home() / 'Library/Mail/V10/MailData/Envelope Index'
MAIL_BASE = Path.home() / 'Library/Mail/V10'
KN = HERE / 'kn.js'

# 兩區設定。STORES：(店碼, 簡報上的店名, 週報 Excel 的店名)
REGIONS = {
    '北一區': dict(
        gen_dir=Path.home() / 'Desktop' / 'north1-weekly-report',
        deck_dir=Path.home() / 'Desktop' / '週報', deck_prefix='北一區週報',
        out_base=Path.home() / '工具中心' / '輸出' / '北一區週報簡報',
        label_fmt='Q{q}W{wk:02d}', dept='01', dss_subject='北一區門市週報_',
        stores=[('004', '士林', '士林門市'), ('005', '微風', '微風門市'), ('024', '美麗華', '美麗華門市'),
                ('046', '阿波羅', '阿波羅門市'), ('054', '高島屋', '大葉高島屋門市'), ('057', '羅東', '羅東門市')],
        alias={},
        repair={'士林維修': '110', '羅東維修': '121'},
        headcount='schedule',
        mysetup={'士林': '士林', '微風': '微風', '美麗華': '美麗華', '阿波羅': '阿波羅', '高島屋': '高島屋'},
        examples=HERE / '評語範例', learn_log=HERE / '評語修改紀錄.md'),
    '北二區': dict(
        gen_dir=Path.home() / 'Desktop' / 'north2-weekly-report',
        deck_dir=Path.home() / 'Desktop' / '週報' / '北二區', deck_prefix='北二區週報',
        out_base=Path.home() / '工具中心' / '輸出' / '北二區週報簡報',
        label_fmt='Q{q}W{wk}', dept='02', dss_subject='北二區門市週報_',
        stores=[('009', '永和', '永和門市'), ('025', '板橋誠品', '板橋誠品門市'), ('050', '西門', '西門門市'),
                ('055', '花蓮', '花蓮門市'), ('063', '板橋遠百', '板橋遠百門市'), ('064', '新莊宏匯', '新莊宏匯門市'),
                ('068', '裕隆城', '新店裕隆城門市')],
        alias={'新西門': '西門', '新店裕隆城': '裕隆城'},
        repair={'花蓮維修': '113'},
        headcount='epb',          # 排班系統帳號看不到北二區 → 用本週有銷售的人數，標黃底請使用者確認
        # 鍵必須跟 warranty_report_all 的店名一致（scan_mysetup 的快取以店名存，鍵不同會漏掉快取日期）
        mysetup={'西門': '西門', '板橋遠百': '板橋遠百', '新莊宏匯': '宏匯', '新店裕隆城': '裕隆城'},
        examples=HERE / '評語範例' / '北二區', learn_log=HERE / '評語修改紀錄_北二區.md',
        # 報告人 Yuni 每週寄出定稿：當下週範本、評語學她的寫法（finals_mail.py）
        final_mail=dict(sender='yuni.wang@studioa.com.tw', subject=r'北二區 (\d\d)Q(\d)W(\d+) 週報'),
        final_dir=Path.home() / '工具中心' / '輸出' / '北二區週報簡報' / '定稿信件'),
}
R = {}


def configure(region):
    """設定本次執行的區域（一次執行只處理一區；模組層級變數供各函式使用）。"""
    global R, REGION, GEN_DIR, DECK_DIR, OUT_BASE, STORES, NAMES
    R = REGIONS[region]
    REGION = region
    GEN_DIR, DECK_DIR, OUT_BASE = R['gen_dir'], R['deck_dir'], R['out_base']
    STORES = R['stores']
    NAMES = [x[1] for x in STORES]
    sys.path.insert(0, str(GEN_DIR))      # 週報產生器 server / multistore_engine 以本區為準


YELLOW = [65535, 62000, 30000]
MONTH_ZH = ['', '一月', '二月', '三月', '四月', '五月', '六月', '七月', '八月', '九月', '十月', '十一月', '十二月']
DEVICES = [('CPU', 'CPU'), ('iPhone', 'iPhone'), ('iPad', 'iPad'), ('Watch', 'Watch'), ('Airpods', 'AirPods')]

WARN = []          # 交付時要提醒使用者的事


def log(msg):
    print(f'[{datetime.now():%H:%M:%S}] {msg}', flush=True)


def warn(msg):
    WARN.append(msg)
    log('⚠️ ' + msg)


def store_key(label):
    """簡報各表的店名寫法不一（大葉 高島屋 / 大葉\n高島屋 / 高島屋門市），統一成 NAMES。"""
    t = re.sub(r'\s', '', str(label)).replace('大葉', '').replace('門市', '')
    t = R.get('alias', {}).get(t, t)
    return t if t in NAMES else None


# ═════════════════════════ 期間 ═════════════════════════
def last_saturday(today):
    return today - timedelta(days=(today.weekday() - 5) % 7 or 7)


class Period:
    def __init__(self, wk_end):
        import multistore_engine as eng
        self.wk_end = wk_end
        self.wk_start = wk_end - timedelta(days=6)
        self.pw_end = self.wk_start - timedelta(days=1)
        self.pw_start = self.pw_end - timedelta(days=6)
        self.m_start = wk_end.replace(day=1)
        self.days_in_month = calendar.monthrange(wk_end.year, wk_end.month)[1]
        fy, q, wk = eng.fy_week(wk_end)
        self.fy, self.q, self.wk = fy, q, wk
        self.label = R['label_fmt'].format(q=q, wk=wk)   # 北一 Q1W03／北二 Q1W3
        self.qtag = f'{fy % 100}Q{q}'                 # 27Q1
        self.full = f'{fy % 100}Q{q}W{wk:02d}'        # 27Q1W03（資料夾、檔名用）


# ═════════════════════════ 週報產生器 Excel ═════════════════════════
def weekly_xlsx(P, out_dir, reuse=None):
    if reuse:
        log(f'沿用週報 Excel：{reuse}')
        return Path(reuse)
    p = out_dir / f'{REGION}週報_{P.wk_end}.xlsx'
    if p.exists():
        log(f'週報 Excel 已存在，沿用：{p.name}')
        return p
    # 使用者每週自己用週報產生器產出、放在下載資料夾（例：北一區週報_2026-10-17.xlsx），有就直接用
    # 只認標準檔名（瀏覽器重複下載的「 (1)」可），不認「_自訂」等自訂期間版本
    pat = re.compile(re.escape(f'{REGION}週報_{P.wk_end}') + r'( \(\d+\))?\.xlsx$')
    hits = sorted((f for f in Path.home().joinpath('Downloads').glob(f'{REGION}週報_{P.wk_end}*.xlsx') if pat.match(f.name)),
                  key=os.path.getmtime)
    if hits:
        shutil.copy(hits[-1], p)
        log(f'週報 Excel（下載資料夾）：{hits[-1].name}')
        return p
    log('下載資料夾沒有本週週報 Excel，自行產生（週報產生器，含 EPB 與 ShopperTrak 人流）…')
    cwd = os.getcwd()
    os.chdir(GEN_DIR)
    try:
        import server
        msgs = []
        data = server._fill_workbook(P.wk_end, lambda m: (msgs.append(m), log('  產生器：' + m)))
    finally:
        os.chdir(cwd)
    for m in msgs:
        if '略過' in m or '失敗' in m or '⚠' in m:
            warn(f'週報產生器：{m}')
    p.write_bytes(data)
    log(f'週報 Excel 完成：{p.name}')
    return p


def recalc_copy(src, out_dir):
    """用 LibreOffice 轉存一份（openpyxl 寫的公式沒有快取值，轉存時會算出來），供讀取比率／合計。"""
    soffice = '/Applications/LibreOffice.app/Contents/MacOS/soffice'
    d = Path(tempfile.mkdtemp(dir=out_dir))
    subprocess.run([soffice, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', str(d), str(src)],
                   capture_output=True, timeout=300)
    got = d / Path(src).name
    if not got.exists():
        raise SystemExit('❌ LibreOffice 重算週報 Excel 失敗（羅東人流等公式值會讀不到）')
    return got


def find_row(ws, text, col=1, start=1):
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, col).value
        if v is not None and str(v).strip().startswith(text):
            return r
    raise KeyError(f'{ws.title} 找不到「{text}」')


def store_rows(ws, start, col=1):
    """從 start 列往下找 6 店的列號 {店名: 列}。"""
    out = {}
    for r in range(start, start + 12):
        k = store_key(ws.cell(r, col).value or '')
        if k and k not in out:
            out[k] = r
    if len(out) != len(NAMES):
        raise KeyError(f'{ws.title} 第 {start} 列起找不到全部 {len(NAMES)} 店：{out}')
    return out


def read_weekly(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    X = {}
    # 重點項目進度：AC+ 佔比（季累積）＝ AC+ 台數 ÷ 主機台數，自己算不吃公式
    ws = wb['重點項目進度']
    ra, rh = find_row(ws, 'AC+ 台數'), find_row(ws, '主機台數')
    hdr = {store_key(ws.cell(ra, c).value): c for c in range(2, len(NAMES) + 3) if store_key(ws.cell(ra, c).value)}
    ac = {}
    for k, dev in enumerate(['CPU', 'iPad', 'iPhone', 'Watch', 'AirPods']):
        assert ws.cell(ra + 1 + k, 1).value == dev and ws.cell(rh + 1 + k, 1).value == dev
        num = {s: ws.cell(ra + 1 + k, c).value or 0 for s, c in hdr.items()}
        den = {s: ws.cell(rh + 1 + k, c).value or 0 for s, c in hdr.items()}
        row = {s: (num[s] / den[s] if den[s] else 0) for s in NAMES}
        row[REGION] = sum(num.values()) / sum(den.values()) if sum(den.values()) else 0
        ac[dev] = row
    X['ac_share'] = ac
    X['ac_title'] = ws.cell(1, 1).value
    rp = find_row(ws, 'ARpedia')
    hdr = {store_key(ws.cell(rp, c).value): c for c in range(2, len(NAMES) + 3) if store_key(ws.cell(rp, c).value)}
    X['arpedia'] = {lab: {s: ws.cell(rp + 1 + k, c).value or 0 for s, c in hdr.items()}
                    for k, lab in enumerate(['本週銷售', '本月銷售', '年度銷售'])}

    # 本週其他細項：主機總數、服務費（P 欄起的子表）
    ws = wb['BY店 本週其他細項']
    rows = store_rows(ws, 3, col=16)
    X['hosts_week'] = {s: ws.cell(r, 17).value or 0 for s, r in rows.items()}
    X['service_fee'] = {s: (ws.cell(r, 20).value or 0) + (ws.cell(r, 21).value or 0) for s, r in rows.items()}

    # 來客數：本週比較／月累積 的人流欄（各店相加）
    def traffic(sheet):
        """台數區塊表頭列有兩個「人流」：第一個是上期、第二個是本期。回傳 (本期, 上期) 6 店合計。"""
        ws = wb[sheet]
        for r in range(1, 40):
            vals = [ws.cell(r, c).value for c in range(1, 30)]
            if '人流' in vals:
                cols = [i + 1 for i, v in enumerate(vals) if v == '人流']
                rows = store_rows(ws, r + 1)
                prev = sum(ws.cell(rr, cols[0]).value or 0 for rr in rows.values())
                cur = sum(ws.cell(rr, cols[1]).value or 0 for rr in rows.values())
                return cur, prev
        raise KeyError(f'{sheet} 找不到人流欄')
    X['traffic_week'], X['traffic_pweek'] = traffic('BY店 本週比較')
    X['traffic_mtd'], X['traffic_pm_same'] = traffic('BY店 月累積')
    return X


# ═════════════════════════ 門市週報 Excel（第2頁業績總表；北一區由 dss@ 寄信，北二區手動下載，同一個工具產生） ═════════════════════════
def _mail_db():
    tmp = Path(tempfile.mkdtemp()) / 'env.db'
    shutil.copy(MAIL_DB, tmp)
    return sqlite3.connect(tmp)


def _emlx_path(rowid):
    hits = glob.glob(str(MAIL_BASE / '**' / 'Messages' / f'{rowid}.emlx'), recursive=True) + \
           glob.glob(str(MAIL_BASE / '**' / 'Messages' / f'{rowid}.partial.emlx'), recursive=True)
    return Path(hits[0]) if hits else None


def mail_xlsx_attachments(rowid):
    """回傳該信所有 xlsx 附件的 bytes（內嵌或 Mail 另存在 Attachments/ 的都找）。"""
    out = []
    p = _emlx_path(rowid)
    if p:
        raw = p.read_bytes()
        raw = raw[raw.index(b'\n') + 1:]
        m = email.message_from_bytes(raw, policy=policy.default)
        for part in m.walk():
            fn = part.get_filename() or ''
            if fn.lower().endswith('.xlsx'):
                b = part.get_payload(decode=True)
                if b and b[:2] == b'PK':
                    out.append(b)
        if not out:
            for f in glob.glob(str(p.parent.parent / 'Attachments' / str(rowid) / '**' / '*.xlsx'), recursive=True):
                out.append(Path(f).read_bytes())
    return out


def dss_store_report(P):
    subj = f"{R['dss_subject']}{P.wk_end + timedelta(days=1):%Y%m%d}"
    for attempt in range(19):                    # 找不到就在背景開 Mail 收信，最多等約 3 分鐘
        wb = _find_dss(subj) or _find_dss_downloads(P)
        if wb:
            return wb
        if attempt == 0:
            log('信箱裡還沒有門市週報 Excel，背景開啟 Mail 收信…')
            subprocess.run(['open', '-g', '-a', 'Mail'])
        import time
        time.sleep(10)
    raise SystemExit(f'❌ 找不到 dss@ 寄的「{subj}」或其附件（通常週日 04:05 寄出）。'
                     f'請確認 Mail 已收到這封信，或把手動產生的「{REGION}_{REGION}週報_YYYYMMDD.xlsx」放在下載資料夾後重跑。')


def _find_dss_downloads(P):
    """使用者手動產生、下載的「北二區_北二區週報_YYYYMMDD.xlsx」。"""
    tag = f'{P.wk_end + timedelta(days=1):%Y%m%d}'
    hits = sorted(Path.home().joinpath('Downloads').glob(f'{REGION}_{REGION}週報_{tag}*.xlsx'), key=os.path.getmtime)
    if hits:
        log(f'門市週報 Excel（下載資料夾）：{hits[-1].name}')
        return openpyxl.load_workbook(hits[-1], data_only=False)
    return None


def _find_dss(subj):
    con = _mail_db()
    rows = con.execute(
        "select m.ROWID from messages m join subjects s on m.subject=s.ROWID "
        "join addresses a on m.sender=a.ROWID where s.subject=? and a.address='dss@studioa.com.tw' "
        "order by m.date_received desc", (subj,)).fetchall()
    for (rid,) in rows:
        for b in mail_xlsx_attachments(rid):
            log(f'門市週報 Excel（信件）：{subj}')
            return openpyxl.load_workbook(BytesIO(b), data_only=False)
    return None


def summary_table(wb, people, off, traffic):
    """照門市週報 Excel『週報』分頁的公式，把 3 個期間（當週/本月累積/上月同期）算出來。
    回傳 {DSS列號: [當週, 本月累積, 本月目標, 達成率, 上月同期, 比]}，None＝該格留空。"""
    sell = wb['sell']
    hdr = [c.value for c in sell[1]][:30]
    S = {}
    for r in range(2, 5):
        row = dict(zip(hdr, [c.value for c in sell[r]][:30]))
        S[row['MAPID']] = {k: (v or 0) for k, v in row.items() if k}
    price = {}
    for r in wb['sacMap'].iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            price[r[0]] = r[1] or 0
    sacJ = defaultdict(float)
    for r in wb['sac'].iter_rows(min_row=2, values_only=True):
        if r[0] in (1, 2, 3):
            sacJ[r[0]] += (r[3] or 0) * price.get(r[1], 0)
    sacK = {i: round(sacJ[i] * 0.5) for i in (1, 2, 3)}
    ttl = {}
    for r in wb['ttl'].iter_rows(min_row=2, values_only=True):
        if r[1] in (1, 2, 3):
            ttl[r[1]] = [x or 0 for x in r[2:6]]
    T = {}
    for r in wb['target'].iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            T[r[0] + 1] = r[2] or 0        # target!C{n} ＝ sort n-1
    cnt = {r[0]: r[1] or 0 for r in wb['sellcnt'].iter_rows(min_row=2, values_only=True) if r[0] in (1, 2, 3)}
    trd = {r[0]: r[1] or 0 for r in wb['tradein'].iter_rows(min_row=2, values_only=True) if r[0] in (1, 2, 3)}

    def v(i, k):
        return S[i].get(k, 0)

    val = {}                                       # 列 → {期間: 值}
    for i in (1, 2, 3):
        a = {}
        a[8] = v(i, 'TTL_AMT')
        a[9] = v(i, 'ttl_amth')
        a[10] = a[8] + a[9]
        a[11] = v(i, 'ttl_amtg')
        a[12] = 0
        a[7] = a[8] + a[11] + sacJ[i]
        a[13] = v(i, 'CPUAMOUNT')
        a[14] = v(i, 'A3PPAMOUNT')
        a[15] = v(i, 'P3PP_AMOUNT_2016')
        a[17] = v(i, 'CC_ACPP') + sacJ[i]
        a[19] = ttl[i][0] + sacK[i]
        a[20] = ttl[i][1]
        a[21] = ttl[i][2]
        a[22] = ttl[i][3] + sacK[i]
        a[23] = v(i, 'CPU_TQTY') + v(i, 'MAC_MINI')
        a[24] = v(i, 'ACPP_CPU') + v(i, 'SAC_CPU')
        a[26] = v(i, 'IPAD_TQTY')
        a[27] = v(i, 'ACPP_IPOD') + v(i, 'SAC_IPAD')
        a[29] = v(i, 'IPHONE_TQTY')
        a[30] = v(i, 'ACPP_IOS') + v(i, 'SAC_IPHONE')
        a[32] = v(i, 'WATCH_TQTY')
        a[33] = v(i, 'ACPP_W') + v(i, 'SAC_WATCH')
        a[35] = v(i, 'air_TQTY')
        a[36] = v(i, 'ACPP_ACC') + v(i, 'SAC_air')
        denom = people if i == 1 else people + off
        a[38] = round(a[8] / denom) if denom else None      # 人數 0（班表抓不到）時留空，不讓整份中斷
        a[39] = traffic[i]
        a[40] = cnt.get(i, 0)
        a[41] = a[40] / a[39] if a[39] else 0
        a[42] = trd.get(i, 0)
        val[i] = a
    tgt = {7: T[2], 8: T[3], 12: T[4], 13: T[5], 14: T[6], 15: T[7], 17: T[24] + T[8], 19: T[9], 20: T[10],
           21: T[11], 22: T[25], 23: T[13] + T[14], 24: T[20] + T[21], 26: T[15], 27: T[26], 29: T[17],
           30: T[27], 32: T[19], 33: T[28], 35: T[23], 36: T[29]}
    ratio_rows = {16: (15, 7), 18: (17, 7), 25: (24, 23), 28: (27, 26), 31: (30, 29), 34: (33, 32), 37: (36, 35)}

    def div(a, b):
        return a / b if b else None

    # 各列的 本月目標／達成率／與上月同期比 照門市週報 Excel 的公式：
    #   一般列 F=D/E、H=(D-G)/G；比率列 F=D/E、H=D-G；
    #   9~11、39~42 沒有目標（10、11 的目標格是常數 0）；12、38 的達成率是 IFERROR(…,0)
    out = {}
    for r in range(7, 43):
        if r in ratio_rows:
            n, d = ratio_rows[r]
            c, m, g = (div(val[i][n], val[i][d]) for i in (1, 2, 3))
            e = div(tgt[n], tgt[d])
            f = div(m, e) if (m is not None and e) else None
            h = (m - g) if (m is not None and g is not None) else None
        else:
            c, m, g = val[1][r], val[2][r], val[3][r]
            e = tgt.get(r, 0 if r in (10, 11) else None)
            f = div(m, e) if e else None
            if r in (12, 38):
                f = div(m, e) if e else 0
            if r == 12:
                g = None
            h = div(m - (g or 0), g) if (g and m is not None) else None
        out[r] = [c, m, e, f, g, h]
    return out


# ═════════════════════════ 排班系統：本週銷售人數 ═════════════════════════
def headcount(P):
    """班表上本週有排任何格子的人（含整週排休；整週空白不算；羅東維修不算）。休假人數固定 0。
    2026-10-13 對 27Q1W02 手工值 44／0 校準。"""
    import schedule_api
    from schedule_compare import STORE_IDS, EXCLUDE_ROLES
    days = [(P.wk_start + timedelta(i)).isoformat() for i in range(7)]
    c = schedule_api.Client()
    total, per = 0, {}
    for store, sid in STORE_IDS.items():
        c.switch_store(sid)
        ps = [p for p in c.periods() if p['start_date'] <= days[-1] and p['end_date'] >= days[0]]
        names = set()
        for p in ps:
            for e in c.schedule(p['id'])['schedule'].values():
                if e.get('role', '') in EXCLUDE_ROLES.get(store, ()):
                    continue
                if any(((e.get('days') or {}).get(d) or {}).get('shift') for d in days):
                    names.add(e.get('name'))
        if not ps:
            warn(f'排班系統：{store} 找不到涵蓋本週的期別')
        per[store] = len(names)
        total += len(names)
    log(f'本週銷售人數 {total}（{per}）')
    return total, 0


def headcount_sales(xlsx_calc):
    """排班系統看不到的區：用週報 Excel「BY店 人員銷售」本週有銷售紀錄的人數。"""
    ws = openpyxl.load_workbook(xlsx_calc, data_only=True)['BY店 人員銷售']
    mode, n = None, 0
    for r in ws.iter_rows(values_only=True):
        if r[1] == '本週':
            mode = 'w'
        elif r[1] == '月累積':
            mode = 'm'
        elif mode == 'w' and r[0] and str(r[0]).startswith('SA'):
            n += 1
    warn(f'本週銷售人數 {n} 是「本週有銷售紀錄的人數」（排班系統無法提供本區人數），已標黃底，請確認後修改（人均產值會跟著變）')
    return n, 0


# ═════════════════════════ 3PP（EPB 重建 DSS） ═════════════════════════
def threepp(start, end):
    import dss_rebuild_from_epb as m
    lines, occ = m.pull(start.isoformat(), end.isoformat(), m.REGIONS[R['dept']])
    d = m.label(lines, occ)
    d['q'] = pd.to_numeric(d.STK_QTY)
    p = d.pivot_table(index='SHOP_ID', columns='type_name', values='q', aggfunc='sum').fillna(0)
    out = {}
    for code, name, _ in STORES:
        out[name] = {dn: tuple(float(p.loc[code].get(f'{k}{t}', 0)) if code in p.index else 0.0
                               for t in ('搭售', '配件', '零搭售')) for k, dn in DEVICES}
    return out


# ═════════════════════════ MySetup ═════════════════════════
def mysetup(P):
    import warranty_report_all as w
    kw = R['mysetup']
    bd = w.scan_mysetup(kw, P.wk_start, P.wk_end)
    got = {date.fromisoformat(k) for k in bd}
    miss = [d for d in (P.wk_start + timedelta(i) for i in range(7)) if d not in got]
    if miss:
        warn('MySetup 缺這幾天的 Personal Setup 信：' + '、'.join(f'{d:%m/%d}' for d in miss) + '（提交數會偏低）')
    s = w.mysetup_sum(bd, P.wk_start, P.wk_end)
    return {store_key(k) or k: s.get(k, 0) for k in kw}     # 快取店名 → 簡報店名（新店裕隆城 → 裕隆城）


# ═════════════════════════ EPB 小查詢 ═════════════════════════
def epb(sql):
    import dss_rebuild_from_epb as m
    return m.epb(sql)


# 第 9 頁「重點項目」每月會換；這裡列有自動算法的品項（本月累積銷售台數，銷退扣回）。
# 鍵＝簡報表格第一欄的文字，值＝EPB poslinev_bi l / stkmas s 的篩選條件。
KEY_ITEMS = {
    'Mac mini M4': "l.cat3_id='3001' and upper(s.name) like '%MAC MINI%M4%'",
}


def key_item_units(P, cond):
    sql = f"""select l.shop_id, sum(case when l.trans_type in ('A','H') then l.stk_qty
                when l.trans_type='E' then -l.stk_qty else 0 end) q
      from poslinev_bi l join stkmas s on s.stk_id=l.stk_id
      where l.org_id='01' and {cond}
        and l.doc_date>=TO_DATE('{P.m_start}','yyyy-mm-dd')
        and l.doc_date<TO_DATE('{P.wk_end + timedelta(days=1)}','yyyy-mm-dd')
        and l.shop_id in ({','.join(repr(s[0]) for s in STORES)})
      group by l.shop_id"""
    df = epb(sql)
    got = {} if df.empty else {r.SHOP_ID: int(float(r.Q)) for r in df.itertuples()}
    return {name: got.get(code, 0) for code, name, _ in STORES}


# ═════════════════════════ 維修中心 ═════════════════════════
def repair(P):
    import repair_report as rr
    periods = rr.calc_periods(P.wk_end)
    res = rr.parse_results(rr.run_epb(rr.build_sql(periods)), periods)
    # 圖表：去年 1 月起每月 ＋ 本月按週（月初~第一個週六為 W1）
    months = []
    y, mth = P.wk_end.year - 1, 1
    while (y, mth) < (P.wk_end.year, P.wk_end.month):
        last = calendar.monthrange(y, mth)[1]
        months.append((f'{mth}月', date(y, mth, 1), date(y, mth, last)))
        y, mth = (y, mth + 1) if mth < 12 else (y + 1, 1)
    weeks, s, k = [], P.m_start, 1
    while s <= P.wk_end:
        e = min(s + timedelta(days=(5 - s.weekday()) % 7), P.wk_end)
        weeks.append((f'W{k}', s, e))
        s, k = e + timedelta(days=1), k + 1
    # 歷史月份沿用 維修圖表歷史.json（使用者原圖的數字，口徑不一、無法重算）；
    # 沒有的完整月份才查 EPB（repair_report 口徑，與週柱、本月累積表一致）並寫回檔案。
    hist_p = HERE / '維修圖表歷史.json'
    hist = json.loads(hist_p.read_text(encoding='utf-8'))
    sids = list(R['repair'].values())
    todo = [(lab, a, b) for lab, a, b in months if any(f'{a:%Y-%m}' not in hist.setdefault(sid, {}) for sid in sids)]
    todo += weeks
    got = {}
    for i in range(0, len(todo), 8):            # 一次查 8 段，避免 SQL 太長
        chunk = {f'{lab}#{a}': (a, b) for lab, a, b in todo[i:i + 8]}
        r = rr.parse_results(rr.run_epb(rr.build_sql(chunk)), chunk)
        for sid in sids:
            for key in chunk:
                got[(sid, key)] = r.get(sid, {}).get(key, {'orig': 0, 'tpp': 0, 'repair': 0})
    series = {}
    for sid in sids:
        for lab, a, b in months:
            ym = f'{a:%Y-%m}'
            if ym not in hist[sid]:
                hist[sid][ym] = got[(sid, f'{lab}#{a}')]
                log(f'維修圖表新增月份 {sid} {ym}：{hist[sid][ym]}')
            series.setdefault(sid, []).append(hist[sid][ym])
        for lab, a, b in weeks:
            series[sid].append(got[(sid, f'{lab}#{a}')])
    hist_p.write_text(json.dumps(hist, ensure_ascii=False, indent=1), encoding='utf-8')
    bars = months + weeks
    labels = [b[0] for b in bars]
    split = 12                                  # 去年 12 個月之後是今年
    return res, labels, split, series


# ═════════════════════════ Keynote ═════════════════════════
def kn(cmd, deck, *args, timeout=900):
    for attempt in range(2):
        r = subprocess.run(['osascript', '-l', 'JavaScript', str(KN), cmd, str(deck), *map(str, args)],
                           capture_output=True, text=True, timeout=timeout)
        # -1728「無法取得物件」：Keynote 還在關上一份同名文件時就去開，偶發；等一下重試一次
        if r.returncode != 0 and '-1728' in r.stderr and attempt == 0:
            import time
            time.sleep(3)
            continue
        break
    if r.returncode != 0:
        raise SystemExit(f'❌ Keynote {cmd} 失敗：{r.stderr.strip()[:800]}')
    return r.stdout


def update_native_charts(deck, jobs, work):
    """維修長條圖保留原生圖表、用 chart_gui 改資料（會操作畫面）。任何一張失敗就全部放棄，改用圖片取代。"""
    if not jobs:
        return
    import chart_gui
    log(f'維修圖表：更新原生圖表資料（{len(jobs)} 張，會操作 Keynote 畫面，請勿動鍵盤滑鼠）…')
    try:
        name = kn('open', deck).strip()
        for j in jobs:
            chart_gui.update_chart(name, j['slide'], j['labels'], j['rows'])
        kn('saveclose', deck)
        log('維修圖表：原生圖表已更新並逐格核對')
        return
    except (Exception, SystemExit) as e:
        try:
            kn('discard', deck)
        except SystemExit:
            pass
        warn(f'維修圖表無法直接更新（{str(e)[:120]}），已改用圖片取代，樣式會與原版略有不同')
    p = work / 'ops_chart_fallback.json'
    p.write_text(json.dumps([op for j in jobs for op in j['fallback']], ensure_ascii=False, indent=1), encoding='utf-8')
    kn('apply', deck, p)


def deck_name(full):
    return f"{R['deck_prefix']} {full}.key"


def find_template(P):
    """本區簡報資料夾裡最新、且早於本週的那份週報 .key。
    有報告人定稿信（北二區 Yuni）時，同一週優先用定稿（含她手改的內容與評語定稿）。"""
    dirs = [DECK_DIR]
    if R.get('final_mail'):
        import finals_mail
        try:
            finals_mail.fetch_finals()
        except Exception as e:
            warn(f'讀取報告人定稿信失敗（{str(e)[:80]}），範本改用自己上週的簡報')
        dirs.append(R['final_dir'])                  # 放後面：同一週時蓋過自己的版本
    best = None
    for d in dirs:
        for f in d.glob(f"{R['deck_prefix']} *.key"):
            m = re.match(re.escape(R['deck_prefix']) + r' (\d\d)Q(\d)W(\d\d)\.key$', f.name)
            if not m:
                continue
            key = (int(m[1]), int(m[2]), int(m[3]))
            if key < (P.fy % 100, P.q, P.wk) and (best is None or key >= best[0]):
                best = (key, f)
    if not best:
        raise SystemExit(f"❌ {DECK_DIR} 找不到上一週的「{R['deck_prefix']} XXQ#W##.key」當範本")
    if R.get('final_mail') and best[1].parent != R['final_dir']:
        warn(f'還沒收到報告人寄出的 {best[1].stem} 定稿信，範本用自己上週的簡報，本週不學評語寫法')
    return best[1]


def src_tables(s):
    return s['tables']


class Deck:
    """dump 出來的結構 + 依內容找表格的小工具。"""

    def __init__(self, dump):
        self.all_slides = dump                                   # 含略過的頁
        self.slides = [s for s in dump if not s.get('skipped')]  # 顯示中的頁

    def tables(self, pred, include_skipped=False):
        src = self.all_slides if include_skipped else self.slides
        return [(s, t) for s in src for t in src_tables(s) if pred(t, s)]

    def opt(self, pred, what):
        hits = self.tables(pred)
        if len(hits) == 1:
            return hits[0]
        warn(f'簡報裡找不到唯一的「{what}」表格（找到 {len(hits)} 個），這一塊略過未更新')
        return None, None

    def one(self, pred, what):
        hits = self.tables(pred)
        if len(hits) != 1:
            raise SystemExit(f'❌ 簡報裡找不到唯一的「{what}」表格（找到 {len(hits)} 個），版面可能改過')
        return hits[0]


def find_3pp(D):
    """3PP 搭售統計表：第二欄是 CPU、26 欄、有 Total 列；北二區頁上另有一張較小的備用表，取列數最多的。"""
    hits = D.tables(lambda t, s: t['cols'] == 26 and cell(t, 0, 1) == 'CPU'
                    and any(cell(t, r, 0) == 'Total' for r in range(t['rows'])))
    if not hits:
        raise SystemExit('❌ 簡報裡找不到 3PP 搭售統計表，版面可能改過')
    return max(hits, key=lambda h: h[1]['rows'])


def cell(t, r, c):
    return t['cells'][r][c][0].strip()


def is_formula(t, r, c):
    return bool(t['cells'][r][c][1])


class Ops(list):
    def set(self, s, t, r, c, v, bg=None, force=False):
        """寫一格；有公式的格子預設不覆蓋（讓 Keynote 自己算），force=True 才寫。"""
        if is_formula(t, r, c) and not force:
            return
        op = {'op': 'cell', 's': s['i'], 't': t['i'], 'r': r, 'c': c}
        if v is not None:
            op['v'] = v
        if bg:
            op['bg'] = bg
        self.append(op)


def num(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer() and abs(v) > 1:
        return int(v)
    return v


RATE2_ROWS = {16, 18}                               # 3PP、保固搭售率：兩位小數
RATE0_ROWS = {25, 28, 31, 34, 37, 41}               # 各品類保固搭售率、提袋率：整數
DASH_ROWS = {10, 11, 12}


def fmt_summary(r, k, v):
    """業績總表比照手做版的顯示格式（門市週報 Excel 的格式）。k：0當週 1本月累積 2目標 3達成率 4上月同期 5比。"""
    if v is None:
        return ''
    if k in (3, 5):                                  # 達成率、與上月同期比
        return f'{v:.2%}'
    if r in RATE2_ROWS:
        return f'{v:.2%}'
    if r in RATE0_ROWS:
        return f'{v:.0%}'
    if v == 0 and r in DASH_ROWS:
        return '-'
    return f'{v:,.0f}'


def chart_image_ops(sl, kind, sid, chart_files):
    """用圖片取代維修長條圖（原生圖表無法更新時的備援；上週已是圖片的就換新圖）。"""
    desc = f'AUTO_CHART_{kind}_{sid}'
    f = chart_files[(kind, sid)]
    if any(m['desc'] == desc for m in sl['images']):
        return [{'op': 'image', 's': sl['i'], 'desc': desc, 'file': str(f)}]
    if not sl['charts']:
        return []
    ops = [{'op': 'image', 's': sl['i'], 'desc': desc, 'chart': 0, 'file': str(f)}]
    for x in sl['texts']:                       # 原圖表上疊的年份字樣、分隔虛線改畫在圖裡
        if re.fullmatch(r'\d{4}年', x['t']):
            ops.append({'op': 'deltext', 's': sl['i'], 'old': x['t']})
    for ln in sl['lines']:
        ops.append({'op': 'delline', 's': sl['i'], 'x': ln['pos']['x'], 'y': ln['pos']['y']})
    return ops


def build_ops(P, D, data, chart_files):
    ops = Ops()
    data['chart_jobs'] = []
    # ── 週次、日期標題 ──
    for s in D.slides:
        for x in s['texts']:
            t = x['t']
            nt = re.sub(r'Q\dW\d{1,2}', P.label, t)
            # 北二區寫「27 Q1W2」「2027 Q1W2」：週次前的年度跟著換（財年跨年時才會變）
            nt = re.sub(r'(?<!\d)(20)?\d\d(?= ?Q\dW)', lambda m: (str(P.fy) if m[1] else f'{P.fy % 100}'), nt)
            if re.fullmatch(r'\d{4}年\d{1,2}/\d{1,2}~\d{4}年\d{1,2}/\d{1,2}週報', t):
                nt = (f'{P.wk_start.year}年{P.wk_start.month}/{P.wk_start.day}~'
                      f'{P.wk_end.year}年{P.wk_end.month}/{P.wk_end.day}週報')
            if nt != t and '北二' not in t:
                ops.append({'op': 'text', 's': s['i'], 'old': t, 'v': nt})

    # ── 第2頁 業績總表 ──
    s, t = D.one(lambda t, s: cell(t, 0, 0).startswith('門市別'), '業績總表')
    people, off = data['people']
    ops.set(s, t, 0, 3, people, bg=YELLOW if (R['headcount'] != 'schedule' or R.get('headcount_fallback')) else None)
    ops.set(s, t, 0, 5, off)
    ops.set(s, t, 1, 0, f'Date : {P.wk_start:%m/%d}～{P.wk_end:%m/%d}')
    ops.set(s, t, 1, 6, P.days_in_month)
    ops.set(s, t, 2, 6, f'{P.wk_end.day / P.days_in_month:.2%}')
    for r, vals in data['summary'].items():
        for k, v in enumerate(vals):
            # 兩區一律照門市週報 Excel 的算法（2026-10-19 使用者指定統一用北一區算法）：
            # 達成率＝本月累積÷目標；與上月同期比：金額/台數列＝(本月−上月)÷上月、搭售率列＝本月−上月。
            # 北二區範本這兩欄是 Keynote 公式（且搭售率列用相對變化），一律用精確值覆蓋。
            ops.set(s, t, r - 3, k + 1, fmt_summary(r, k, v), force=k in (3, 5))

    # ── 重點項目進度 ──
    s, t = D.opt(lambda t, s: re.fullmatch(r'\d\dQ\d', cell(t, 0, 0) or '') and cell(t, 1, 0) == '重點項目', 'AC+ 佔比')
    for c in range(t['cols'] if t else 0):
        if re.fullmatch(r'\d\dQ\d', cell(t, 0, c)):
            ops.set(s, t, 0, c, P.qtag)
    cols = {store_key(cell(t, 1, c)) or cell(t, 1, c): c for c in range(1, t['cols'])} if t else {}
    for r in range(3, t['rows'] if t else 0):
        dev = cell(t, r, 0)
        for name, c in cols.items():
            if dev in data['ac_share'] and name in data['ac_share'][dev]:
                ops.set(s, t, r, c, data['ac_share'][dev][name])
    s, t = D.one(lambda t, s: cell(t, 0, 0) == 'ARpedia', 'ARpedia')
    cols = {store_key(cell(t, 0, c)): c for c in range(1, t['cols']) if store_key(cell(t, 0, c))}
    for r in range(1, 4):
        lab = cell(t, r, 0)
        for name, c in cols.items():
            ops.set(s, t, r, c, data['arpedia'][lab][name])
    hdr = [cell(t, 0, c) for c in range(t['cols'])]
    k = hdr.index('應達成率')
    ops.set(s, t, 2, k, P.wk_end.day / P.days_in_month, force=True)
    ops.set(s, t, 3, k, P.wk_end.timetuple().tm_yday / (366 if calendar.isleap(P.wk_end.year) else 365), force=True)
    s, t = D.opt(lambda t, s: cell(t, 0, 0) == '重點項目' and store_key(cell(t, 0, 1)), '重點項目（本月）')
    if t:
        cols = {store_key(cell(t, 0, c)): c for c in range(1, t['cols']) if store_key(cell(t, 0, c))}
        hdr = [cell(t, 0, c) for c in range(t['cols'])]
        for r in range(1, t['rows']):
            item = cell(t, r, 0)
            got = data['key_items'].get(item)
            for name, c in cols.items():
                ops.set(s, t, r, c, got[name] if got else '', bg=None if got else YELLOW)
            if '應達成率' in hdr:
                ops.set(s, t, r, hdr.index('應達成率'), P.wk_end.day / P.days_in_month, force=True)
            if not got:
                warn(f'重點項目「{item}」沒有自動算法，已標黃底請手填（要自動化請在 KEY_ITEMS 加條件）')

    # ── MySetup ──
    s, t = D.one(lambda t, s: 'Total 銷售台數' in cell(t, 0, 1), 'MySetup')
    for r in range(1, t['rows']):
        name = store_key(cell(t, r, 0))
        if name:
            ops.set(s, t, r, 1, data['hosts_week'][name])
            ops.set(s, t, r, 2, data['mysetup'][name])
            ops.set(s, t, r, 4, data['service_fee'][name])

    # ── 3PP 搭售統計 ──
    s, t = find_3pp(D)
    dev_col = {cell(t, 0, c): c for c in range(t['cols']) if cell(t, 0, c) in [d for _, d in DEVICES]}

    def write_3pp_row(r, vals, force=False):
        for dev, c in dev_col.items():
            att, acc, zero = vals[dev]
            ops.set(s, t, r, c, num(att), force=force)
            ops.set(s, t, r, c + 1, num(acc), force=force)
            ops.set(s, t, r, c + 2, acc / att if att else 0, force=force)
            ops.set(s, t, r, c + 3, num(zero), force=force)
            ops.set(s, t, r, c + 4, zero / (att + zero) if (att + zero) else 0, force=force)
    total_r = next(r for r in range(t['rows']) if cell(t, r, 0) == 'Total')
    # 北二區的主表是公式，引用頁外一張「貼 DSS 匯出值」的來源表（列順序＝DSS 門市代碼順序，
    # 左欄店名沒維護、比率存成 300 而非 300%）→ 寫進來源表，主表自己會帶出來
    m_src = re.match(r"='?([^':]+)'?::", t['cells'][1][1][1] or '')
    src = next((x for x in s['tables'] if m_src and x['name'] == m_src[1]), None)
    if src:
        log(f"3PP 主表引用「{src['name']}」，改寫入來源表（依門市代碼順序）")
        for k, (_, name, _) in enumerate(STORES):
            for dev, c in dev_col.items():
                att, acc, zero = data['3pp_week'][name][dev]
                for off_, v in enumerate([num(att), num(acc), round(acc / att * 100) if att else 0, num(zero),
                                          round(zero / (att + zero) * 100) if (att + zero) else 0]):
                    ops.append({'op': 'cell', 's': s['i'], 't': src['i'], 'r': k + 1, 'c': c + off_, 'v': v})
    else:
        for r in range(1, total_r):
            name = store_key(cell(t, r, 0))
            if name:
                write_3pp_row(r, data['3pp_week'][name])
    need = MONTH_ZH[(P.wk_end.month - 2) % 12 + 1]
    first_month = cell(t, total_r + 1, 0).replace('ㄧ', '一')
    if first_month != need:
        # 新的月份：表格加一列，既有月份列整列往下搬一格，空出 Total 下方那列寫上個月整月。
        # 搬的時候只搬台數，比率由台數重算（顯示值是四捨五入過的，不能拿來搬）。
        def parse(v):
            v = v.replace(',', '').replace('(', '-').replace(')', '').strip()
            try:
                return float(v)
            except ValueError:
                return 0.0
        months = []
        for r in range(total_r + 1, t['rows']):
            lab = cell(t, r, 0)
            if not lab:
                continue
            months.append((lab, {dev: tuple(parse(cell(t, r, c + k)) for k in (0, 1, 3)) for dev, c in dev_col.items()}))
        ops.append({'op': 'rowcount', 's': s['i'], 't': t['i'], 'n': t['rows'] + 1})
        mv = {dev: tuple(sum(data['3pp_month'][n][dev][k] for n in NAMES) for k in range(3)) for dev in dev_col}
        months.insert(0, (need, mv))

        def put(r, lab, vals):
            ops.append({'op': 'cell', 's': s['i'], 't': t['i'], 'r': r, 'c': 0, 'v': lab})
            for dev, c in dev_col.items():
                att, acc, zero = vals[dev]
                for off_, v in enumerate([num(att), num(acc), acc / att if att else 0, num(zero),
                                          zero / (att + zero) if (att + zero) else 0]):
                    ops.append({'op': 'cell', 's': s['i'], 't': t['i'], 'r': r, 'c': c + off_, 'v': v})
        for k, (lab, vals) in enumerate(months):
            put(total_r + 1 + k, lab, vals)
        log(f'3PP 表新增月份列：{need}')

    # ── 維修中心預約（手填：清空＋黃底） ──
    for key in ('STUDIO A', 'Apple'):
        s, t = D.one(lambda t, s, key=key: cell(t, 0, 0) == key and cell(t, 0, 1) == '可預約數', f'維修預約 {key}')
        for r in range(1, t['rows']):
            for c in (1, 2, 4):
                ops.set(s, t, r, c, '', bg=YELLOW)

    # ── 維修中心：營收表、圖表、SACare 整機（手填） ──
    res, labels, split = data['repair'][:3]
    store_of_slide, cur = {}, None
    for sl in D.all_slides:
        for x in sl['texts']:
            if x['t'] in R['repair']:
                cur = R['repair'][x['t']]
        store_of_slide[sl['i']] = cur
    rk = {'保外維修': 'repair', '原廠周邊': 'orig', '3PP': 'tpp'}
    plab = ['本月累積', '上月同期', '去年同期', '今年累積', '去年累積']
    for sl in D.all_slides:
        sid = store_of_slide[sl['i']]
        for t in sl['tables']:
            a = cell(t, 0, 0)
            if a in rk and t['rows'] >= 6 and cell(t, 1, 0) == '本月累積' and sid:
                for r in range(1, 6):
                    assert cell(t, r, 0) == plab[r - 1]
                    ops.set(sl, t, r, 1, res.get(sid, {}).get(plab[r - 1], {}).get(rk[a], 0))
            if a == 'SACare 整機搭售' and not sl.get('skipped'):
                for r in (1, 2):
                    for c in (1, 2, 3):
                        ops.set(sl, t, r, c, '', bg=YELLOW)
        title = ' '.join(x['t'] for x in sl['texts'])
        kind = '3pp' if '3PP營業收入' in title else ('repair' if '保外維修收入' in title else None)
        if kind and sid and not sl.get('skipped'):
            if sl['charts']:
                # 保留原生圖表（樣式與原版完全相同）：資料由 chart_gui 操作圖表資料編輯器寫入，
                # 失敗時改用 fallback（圖片取代）
                c = sl['charts'][0]
                ops.append({'op': 'yearmark', 's': sl['i'], 'x': round(c['pos']['x'] + c['w'] * split / len(labels)),
                            'prev': f'{P.wk_end.year - 1}年', 'cur': f'{P.wk_end.year}年'})
                keys = ('orig', 'tpp') if kind == '3pp' else ('repair',)
                data['chart_jobs'].append({'slide': sl['i'] + 1, 'labels': labels,
                                           'rows': [[round(x[k]) for x in data['repair'][3][sid]] for k in keys],
                                           'fallback': chart_image_ops(sl, kind, sid, chart_files)})
            else:
                ops.extend(chart_image_ops(sl, kind, sid, chart_files))

    # ── CSAT（手填：黃底，保留上週文字當格式參考） ──
    for s, t in D.tables(lambda t, s: '店點狀況' in cell(t, 0, 1)):
        for r in range(1, t['rows']):
            if '滿意度' in cell(t, r, 0):
                ops.set(s, t, r, 1, None, bg=YELLOW)
                ops.set(s, t, r, 2, None, bg=YELLOW)
    return ops


def make_charts(P, data, out_dir):
    import charts
    res, labels, split, series = data['repair']
    files = {}
    D = data['deck_dump']
    geo = {}
    for sl in D.slides:
        title = ' '.join(x['t'] for x in sl['texts'])
        sid = next((v for k, v in R['repair'].items() if k in title), None)
        kind = '3pp' if '3PP營業收入' in title else ('repair' if '保外維修收入' in title else None)
        if sid and kind:
            box = (sl['images'] and next((m for m in sl['images'] if m['desc'] == f'AUTO_CHART_{kind}_{sid}'), None)) \
                  or (sl['charts'][0] if sl['charts'] else None)
            geo[(kind, sid)] = (box['w'], box['h'])
    years = (f'{P.wk_end.year - 1}年', f'{P.wk_end.year}年', split)
    for (kind, sid), (w, h) in geo.items():
        f = out_dir / f'chart_{kind}_{sid}.png'
        s = series[sid]
        if kind == '3pp':
            charts.render(f, labels, years, [('原廠周邊', [x['orig'] for x in s], charts.BLUE),
                                             ('3PP', [x['tpp'] for x in s], charts.GREEN)], w, h, stacked=True)
        else:
            charts.render(f, labels, years, [('保外維修', [x['repair'] for x in s], charts.BLUE)], w, h, stacked=False)
        files[(kind, sid)] = f
    return files


# ═════════════════════════ 評語素材 ═════════════════════════
def sheet_md(ws, max_col=40, max_row=None):
    lines = []
    for row in ws.iter_rows(max_row=max_row, max_col=max_col, values_only=True):
        vals = ['' if v is None else (f'{v:.4g}' if isinstance(v, float) else str(v)) for v in row]
        while vals and vals[-1] == '':
            vals.pop()
        if vals:
            lines.append(' | '.join(v.replace('\n', ' ') for v in vals))
    return '\n'.join(lines)


def comment_material(P, data, xlsx_calc, out_dir, prev_comments):
    wb = openpyxl.load_workbook(xlsx_calc, data_only=True)
    md = [f'# {REGION}週報 {P.full} 評語素材（本週 {P.wk_start:%m/%d}~{P.wk_end:%m/%d}，上週 {P.pw_start:%m/%d}~{P.pw_end:%m/%d}）\n']
    md.append('## 上週簡報的評語（追蹤「先前要求的事項」用）\n')
    for store, (a, b) in prev_comments.items():
        md.append(f'### {store}\n- 上週店點狀況：{a}\n- 執行進度與改善狀況：{b}\n')
    prev = Period(P.wk_end - timedelta(days=7))
    pf = OUT_BASE / prev.full / 'facts.json'
    if pf.exists():
        pfacts = json.loads(pf.read_text(encoding='utf-8'))
        md.append('## 上週快照（上週產出時存的數字）\n')
        md.append('ARpedia 上週銷售：' + ' | '.join(f'{n} {pfacts["arpedia_week"].get(n, "-")}' for n in NAMES))
        md.append('MySetup 上週提交：' + ' | '.join(f'{n} {pfacts["mysetup"].get(n, "-")}' for n in NAMES) + '\n')
    (out_dir / 'facts.json').write_text(json.dumps({'arpedia_week': data['arpedia']['本週銷售'],
                                                    'mysetup': data['mysetup']}, ensure_ascii=False), encoding='utf-8')
    md.append('## 3PP 搭售／零搭售（台數；配件＝件數）\n')
    md.append('門市 | 品類 | 本週 搭售/配件/零搭售 | 上週 搭售/配件/零搭售')
    for n in NAMES:
        for _, dev in DEVICES:
            w, pw = data['3pp_week'][n][dev], data['3pp_pweek'][n][dev]
            md.append(f'{n} | {dev} | {w[0]:.0f}/{w[1]:.0f}/{w[2]:.0f} | {pw[0]:.0f}/{pw[1]:.0f}/{pw[2]:.0f}')
    md.append('\n## MySetup 本週提交數 ／ 主機總數\n')
    md.append(' | '.join(f'{n} {data["mysetup"].get(n, "-")}/{data["hosts_week"][n]}' for n in NAMES))
    for name in ['BY店 本週比較', 'BY店 本週其他細項', '重點項目進度', 'BY店 月累積']:
        md.append(f'\n## {name}\n')
        md.append(sheet_md(wb[name]))
    for ws in wb.worksheets:
        if ws.title.startswith('配件'):
            md.append(f'\n## {ws.title}（只列 週對比 欄）\n')
            md.append(sheet_md(ws, max_col=6))
    p = out_dir / '評語素材.md'
    p.write_text('\n'.join(md), encoding='utf-8')
    return p


def learn_from_last_week(tmpl, D):
    """上週簡報裡的評語＝使用者改完的定稿。存進 評語範例/，並和上週自動初稿（評語.json）比對，
    有改的地方記進 評語修改紀錄.md —— 下次寫評語前先讀，逐週貼近使用者的寫法。"""
    m = re.match(re.escape(R['deck_prefix']) + r' (\d\dQ\dW\d\d)\.key$', tmpl.name)
    final = prev_comments_of(D)
    if not m or not final:                       # 範本不是標準檔名（例如測試）就不學
        return
    prev_full = m[1]
    if R.get('final_mail') and tmpl.parent != R['final_dir']:     # 只有報告人寄回的定稿才算定稿
        return
    R['examples'].mkdir(parents=True, exist_ok=True)
    ex = R['examples'] / f"{R['deck_prefix']} {prev_full}.txt"
    ex.write_text(''.join(f'【{k}】\nA: {a}\nB: {b}\n' for k, (a, b) in final.items()), encoding='utf-8')
    draft_p = OUT_BASE / prev_full / '評語.json'
    if not draft_p.exists():
        return
    draft = json.loads(draft_p.read_text(encoding='utf-8'))
    log_p = R['learn_log']
    if log_p.exists() and f'## {prev_full}' in log_p.read_text(encoding='utf-8'):
        return
    lines = [f'\n## {prev_full}（初稿 → 使用者定稿）\n']
    changed = 0
    for k, (a, b) in final.items():
        for lab, d, f in (('上週店點狀況', draft.get(k, ['', ''])[0], a), ('執行進度與改善狀況', draft.get(k, ['', ''])[1], b)):
            if d.strip() == f.strip():
                lines.append(f'- {k}／{lab}：未修改（初稿直接採用）')
            else:
                changed += 1
                lines.append(f'- {k}／{lab}\n  - 初稿：{d}\n  - 定稿：{f}')
    with log_p.open('a', encoding='utf-8') as fh:
        fh.write('\n'.join(lines) + '\n')
    log(f'評語學習：{prev_full} 共 {changed} 段被修改，已記入 評語修改紀錄.md')


def prev_comments_of(D):
    out = {}
    for s, t in D.tables(lambda t, s: '店點狀況' in cell(t, 0, 1)):
        for r in range(1, t['rows']):
            k = store_key(cell(t, r, 0))
            if k:
                out[k] = (cell(t, r, 1), cell(t, r, 2))
    return out


# ═════════════════════════ main ═════════════════════════
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--region', default='北一區', choices=list(REGIONS))
    ap.add_argument('--wk-end', type=date.fromisoformat)
    ap.add_argument('--template', type=Path, help='範本 .key（預設：桌面/週報 最新的上一週）')
    ap.add_argument('--out-dir', type=Path, help='簡報輸出資料夾（預設：桌面/週報）')
    ap.add_argument('--xlsx', help='沿用已產生的週報 Excel')
    ap.add_argument('--force', action='store_true', help='目標 .key 已存在也覆蓋')
    a = ap.parse_args()
    configure(a.region)

    P = Period(a.wk_end or last_saturday(date.today()))
    # 指定 --out-dir（測試）時，中間檔與狀態也放在那裡，不覆寫正式輸出資料夾、不寫評語學習紀錄
    work = (a.out_dir / P.full) if a.out_dir else OUT_BASE / P.full
    work.mkdir(parents=True, exist_ok=True)
    log(f'══ {REGION}週報 {P.full}（{P.wk_start} ~ {P.wk_end}）══')

    tmpl = a.template or find_template(P)
    target = (a.out_dir or DECK_DIR) / deck_name(P.full)
    if target.exists() and not a.force:
        raise SystemExit(f'❌ {target} 已存在；確定要重做請加 --force')
    log(f'範本：{tmpl.name} → {target}')

    tmp_deck = work / ('_tmp_' + deck_name(P.full))
    if tmp_deck.exists():
        shutil.rmtree(tmp_deck) if tmp_deck.is_dir() else tmp_deck.unlink()
    (shutil.copytree if tmpl.is_dir() else shutil.copy)(tmpl, tmp_deck)
    D = Deck(json.loads(kn('dump', tmp_deck)))

    data = {'deck_dump': D}
    xlsx = weekly_xlsx(P, work, a.xlsx)
    xlsx_calc = recalc_copy(xlsx, work)          # 羅東人流、合計等是公式，要算過才讀得到
    X = read_weekly(xlsx_calc)
    data.update(X)
    data['people'] = None
    if R['headcount'] == 'schedule':
        try:
            data['people'] = headcount(P)
        except (Exception, SystemExit) as e:           # schedule_api 登入失敗會 sys.exit
            warn(f'排班系統讀取失敗（{str(e)[:80]}），本週銷售人數改用「本週有銷售的人數」')
    if not data['people'] or not data['people'][0]:
        data['people'] = headcount_sales(xlsx_calc)
        R['headcount_fallback'] = True
    dwb = dss_store_report(P)
    data['summary'] = summary_table(dwb, *data['people'],
                                    {1: X['traffic_week'], 2: X['traffic_mtd'], 3: X['traffic_pm_same']})
    log('3PP 重建（本週、上週）…')
    data['3pp_week'] = threepp(P.wk_start, P.wk_end)
    data['3pp_pweek'] = threepp(P.pw_start, P.pw_end)
    _, t3 = find_3pp(D)
    tr = next(r for r in range(t3['rows']) if cell(t3, r, 0) == 'Total')
    if cell(t3, tr + 1, 0).replace('ㄧ', '一') != MONTH_ZH[(P.wk_end.month - 2) % 12 + 1]:
        pm_end = P.m_start - timedelta(days=1)
        log(f'3PP 重建（上個月整月 {pm_end:%Y-%m}，新增月份列用）…')
        data['3pp_month'] = threepp(pm_end.replace(day=1), pm_end)
    data['mysetup'] = mysetup(P)
    data['key_items'] = {}
    for _, tk in D.tables(lambda t, s: cell(t, 0, 0) == '重點項目' and store_key(cell(t, 0, 1))):
        for r in range(1, tk['rows']):
            item = cell(tk, r, 0)
            if item in KEY_ITEMS:
                data['key_items'][item] = key_item_units(P, KEY_ITEMS[item])
    log('維修中心…')
    data['repair'] = repair(P)

    files = make_charts(P, data, work)
    ops = build_ops(P, D, data, files)
    (work / 'ops.json').write_text(json.dumps(ops, ensure_ascii=False, indent=1), encoding='utf-8')
    log(f'寫入 Keynote（{len(ops)} 個操作）…')
    kn('apply', tmp_deck, work / 'ops.json')
    update_native_charts(tmp_deck, data['chart_jobs'], work)
    shutil.move(tmp_deck, target)

    if not a.out_dir:
        learn_from_last_week(tmpl, D)
    mat = comment_material(P, data, xlsx_calc, work, prev_comments_of(D))
    repairs = '、'.join(k.replace('維修', '') for k in R['repair'])
    manual = ['維修中心預約', f'SACare 整機搭售（{repairs} 3PP 頁）', '維修滿意度 CSAT']
    if R['headcount'] != 'schedule' or R.get('headcount_fallback'):
        manual.insert(0, '業績總表「本週銷售人數」（暫填本週有銷售的人數）')
    state = {'region': REGION, 'label': P.full, 'deck': str(target), 'material': str(mat), 'xlsx': str(xlsx),
             'stores': NAMES, 'warnings': WARN, 'manual': manual}
    (work / 'state.json').write_text(json.dumps(state, ensure_ascii=False, indent=1), encoding='utf-8')
    log(f'✅ 簡報數字完成：{target}')
    log(f'   評語素材：{mat}')
    for w_ in WARN:
        log(f'   ⚠️ {w_}')
    print(json.dumps(state, ensure_ascii=False))


if __name__ == '__main__':
    main()
